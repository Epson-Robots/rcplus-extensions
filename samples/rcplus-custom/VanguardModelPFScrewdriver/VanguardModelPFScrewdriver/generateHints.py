#!/usr/bin/env python3
# -*- coding: utf-8 -*-


"""
Developer tool for generating Resources/Hints.json from Hints.xlsx.

Hints.xlsx is the source data.
Resources/Hints.json is the generated file used by the Extension.
Normally, this script does not need to be executed.

Run:
    python generateHints.py

Requires:
    openpyxl
"""


from openpyxl import load_workbook
from json import dump
from pathlib import Path


class HintDataGenerator():
    def __init__(self, input, output):
        self.input = input
        self.output = output
        self.records = list()

    def readInput(self):
        book = load_workbook(self.input, read_only=True, data_only=True)
        sheet = book['Hints']
        headers = list()
        for row in sheet.iter_rows(1, 1):
            for cell in row:
                headers.append(cell.value)
        for row in sheet.iter_rows(2):
            fields = list()
            for cell in row:
                fields.append(cell.value)
            record = dict(zip(headers, fields))
            self.records.append(record)

    def writeOutputJson(self):
        data = {
            'Hints': dict(),
            'References': dict()
        }
        for record in self.records:
            if not record['ID']:
                continue
            if '=' in record['ID']:
                id, referredId = record['ID'].split('=')
                data['References'][id] = referredId
            else:
                id = record['ID']
                if id not in data['Hints']:
                    data['Hints'][id] = dict()
                pattern = f'{record["ToolType"]}_{record["Mode"]}'
                format = record['Format']
                if format.startswith('"'):
                    caption = ""
                    format = format.strip('"')
                else:
                    caption = format
                    format = ''
                data['Hints'][id][pattern] = {
                    'Min': record['Min'],
                    'Max': record['Max'],
                    'Factor': record['Factor'] if record['Factor'] else 0,
                    'Caption':  caption,
                    'Unit': record['Unit'] or '',
                    'Format': format,
                }
        with open(self.output, 'wt', encoding='utf-8') as file:
            dump(data, file, indent=2)

    def run(self):
        self.readInput()
        self.writeOutputJson()


if __name__ == '__main__':
    from argparse import ArgumentParser

    parser = ArgumentParser()
    parser.add_argument(
        '--input',
        help='input Excel path',
        default=Path('Hints.xlsx')
    )
    parser.add_argument(
        '--output',
        help='output JSON path',
        default=(Path('Resources') / "Hints.json")
    )
    arg = parser.parse_args()

    HintDataGenerator(arg.input, arg.output).run()
